#!/usr/bin/env python3
"""Static gate for generated selected task manifests.

This checker intentionally reads task paths from a selected manifest instead of
walking run directories. The stress-generation runs contain many obsolete rerun
artifacts, and scanning them directly produces false positives.
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import tarfile
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ISSUE_KEYS = [
    "missing_upload",
    "fake_media",
    "bad_document_format",
    "bad_archive_member_format",
    "fake_audio_text",
    "old_osmand_path",
    "old_audio_recorder_path",
    "visible_leak",
    "bad_linux_path",
    "single_topology",
    "surface_schema",
    "metadata_trace",
    "clear_before_ensure_setup",
    "development_manifest_upload",
    "unsupported_image_dimension_rule",
    "explicit_guard_scored",
    "android_helper_before_ensure",
    "agent_inaccessible_path",
    "declared_device_unused",
    "xlsx_sheet_missing",
    "shell_expected_transposed",
]


ANDROID_HELPER_APP_CATEGORIES: list[tuple[re.Pattern[str], str]] = [
    (re.compile(r"^androidworld_sms_", re.I), "sms"),
    (re.compile(r"^androidworld_contact", re.I), "contacts"),
    (re.compile(r"^androidworld_calendar", re.I), "calendar"),
    (re.compile(r"^androidworld_(clock|alarm|timer)", re.I), "clock"),
    (re.compile(r"^androidworld_task", re.I), "tasks"),
    (re.compile(r"^androidworld_retro_music|^androidworld_mp3_", re.I), "retro_music"),
    (re.compile(r"^androidworld_osmand|^osmand_", re.I), "osmand"),
    (re.compile(r"^androidworld_recipe|^androidworld_recipes|^broccoli_", re.I), "broccoli"),
]

ANDROID_ENSURE_APP_CATEGORIES: list[tuple[re.Pattern[str], str]] = [
    (re.compile(r"sms", re.I), "sms"),
    (re.compile(r"contact", re.I), "contacts"),
    (re.compile(r"calendar", re.I), "calendar"),
    (re.compile(r"clock", re.I), "clock"),
    (re.compile(r"tasks?", re.I), "tasks"),
    (re.compile(r"retro", re.I), "retro_music"),
    (re.compile(r"osmand", re.I), "osmand"),
    (re.compile(r"broccoli|recipe", re.I), "broccoli"),
]


def _load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text())


def _json_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False)


def _allowed_surface_names(repo_root: Path) -> set[str]:
    relative = Path("mdcbench/tasks/generation/profiles/linux_android/surface_capabilities.yaml")
    candidates = [
        repo_root / relative,
        Path(__file__).resolve().parents[2] / relative,
    ]
    for path in candidates:
        if not path.exists():
            continue
        names: set[str] = set()
        in_surfaces = False
        for line in path.read_text(encoding="utf-8").splitlines():
            if line.strip() == "surfaces:":
                in_surfaces = True
                continue
            if not in_surfaces:
                continue
            match = re.match(r"^  ([a-z]+(?:\.[A-Za-z0-9_-]+)+):\s*$", line)
            if match:
                names.add(match.group(1))
        if names:
            return names
    return set()


def _resolve_repo_path(repo_root: Path, path: str) -> Path:
    return Path(path.replace("${repo_root}", str(repo_root)))


def _file_description(path: Path) -> str:
    result = subprocess.run(
        ["file", "-b", str(path)],
        capture_output=True,
        text=True,
        check=False,
    )
    return result.stdout.lower().strip()


def _media_format_ok(suffix: str, description: str) -> bool:
    if suffix == ".mp3":
        return "mpeg" in description or "audio" in description
    if suffix == ".wav":
        return "wave audio" in description
    if suffix in {".jpg", ".jpeg"}:
        return "jpeg image" in description
    if suffix == ".png":
        return "png image" in description
    if suffix == ".m4a":
        return "audio" in description or "iso media" in description
    return True


def _is_probably_text_asset(path: Path) -> bool:
    return path.suffix.lower() in {
        ".csv",
        ".gpx",
        ".html",
        ".htm",
        ".json",
        ".md",
        ".m3u",
        ".m3u8",
        ".txt",
        ".xml",
        ".eml",
        ".yaml",
        ".yml",
    }


def _development_trace_pattern(*, include_eval_terms: bool = False) -> re.Pattern[str]:
    extra_terms = (
        r"|test harness|automated checks|oracle|evaluator|materializer checks|"
        r"exact description text to use|copy verbatim into|deterministic format"
        if include_eval_terms
        else ""
    )
    return re.compile(
        r"mdcbench|real100|real200|real300|generation_engine_probe|"
        r"stable submit token|submission_key|submit_hash|machine-friendly|"
        r"automated detection|deterministic verification|placeholder media|"
        r"\bcreator=\"fixture\"|fixture/|"
        rf"assembled_task|stage4_assets|sample_\d{{6}}{extra_terms}",
        re.I,
    )


def _asset_text_contains_trace(path: Path) -> tuple[bool, str | None]:
    pattern = _development_trace_pattern(include_eval_terms=True)
    try:
        if _is_probably_text_asset(path):
            text = path.read_text(encoding="utf-8", errors="ignore")
            match = pattern.search(text)
            return bool(match), match.group(0) if match else None
        if path.suffix.lower() == ".zip":
            with zipfile.ZipFile(path) as archive:
                for info in archive.infolist():
                    if info.is_dir():
                        continue
                    inner = Path(info.filename)
                    if not _is_probably_text_asset(inner):
                        continue
                    data = archive.read(info.filename, pwd=None)
                    text = data.decode("utf-8", errors="ignore")
                    match = pattern.search(text)
                    if match:
                        return True, f"{info.filename}:{match.group(0)}"
    except Exception as exc:  # pragma: no cover - diagnostic gate should not crash on one asset.
        return True, f"asset_scan_error:{exc}"
    return False, None


def _strip_local_paths(value: Any) -> Any:
    if isinstance(value, dict):
        return {
            key: ("<local_path>" if key == "local_path" else _strip_local_paths(child))
            for key, child in value.items()
        }
    if isinstance(value, list):
        return [_strip_local_paths(child) for child in value]
    return value


def _setup_visible_trace(task: dict[str, Any]) -> str | None:
    """Scan setup-provided app/user content while ignoring host local_path values."""
    pattern = _development_trace_pattern(include_eval_terms=True)
    visible_ops: list[Any] = []
    visible_setup_types = {
        "androidworld_sms_receive",
        "androidworld_sms_send",
        "androidworld_contact_add",
        "androidworld_calendar_event_add",
        "androidworld_alarm_add",
        "androidworld_timer_add",
        "androidworld_recipe_add",
        "androidworld_task_add",
    }
    for block in task.get("setup", []) or []:
        if not isinstance(block, dict):
            continue
        for op in block.get("config", []) or []:
            if isinstance(op, dict) and op.get("type") in visible_setup_types:
                visible_ops.append(op)
    setup_text = _json_text(_strip_local_paths(visible_ops))
    match = pattern.search(setup_text)
    return match.group(0) if match else None


def _document_format_error(path: Path) -> str | None:
    suffix = path.suffix.lower()
    try:
        if suffix == ".xlsx":
            load_workbook(path, read_only=True).close()
        elif suffix == ".docx":
            with zipfile.ZipFile(path) as archive:
                names = set(archive.namelist())
                required = {"[Content_Types].xml", "_rels/.rels", "word/document.xml"}
                missing = sorted(required - names)
                if missing:
                    return f"missing:{','.join(missing)}"
        elif suffix in {".odt", ".odp", ".ods"}:
            with zipfile.ZipFile(path) as archive:
                names = set(archive.namelist())
                required = {"mimetype", "content.xml", "META-INF/manifest.xml"}
                missing = sorted(required - names)
                if missing:
                    return f"missing:{','.join(missing)}"
                mimetype = archive.read("mimetype").decode("utf-8", errors="ignore")
                expected = {
                    ".odt": "application/vnd.oasis.opendocument.text",
                    ".odp": "application/vnd.oasis.opendocument.presentation",
                    ".ods": "application/vnd.oasis.opendocument.spreadsheet",
                }[suffix]
                if expected not in mimetype:
                    return f"mimetype:{mimetype}"
        elif suffix == ".pdf":
            if not path.read_bytes().startswith(b"%PDF"):
                return "missing_pdf_magic"
    except Exception as exc:
        return str(exc)
    return None


def _member_extension(name: str) -> str:
    lower = name.lower()
    for suffix in (".tar.gz", ".tgz"):
        if lower.endswith(suffix):
            return suffix
    return Path(name).suffix.lower()


def _document_bytes_format_error(name: str, data: bytes) -> str | None:
    suffix = _member_extension(name)
    try:
        if suffix == ".json":
            json.loads(data.decode("utf-8"))
        elif suffix in {".jpg", ".jpeg", ".png", ".mp3", ".wav", ".m4a", ".mp4"}:
            with tempfile.NamedTemporaryFile(suffix=suffix) as tmp:
                tmp.write(data)
                tmp.flush()
                description = _file_description(Path(tmp.name))
            if not _media_format_ok(suffix, description):
                return f"bad_media:{description}"
        elif suffix == ".pdf":
            if not data.startswith(b"%PDF"):
                return "missing_pdf_magic"
        elif suffix in {".tar", ".tar.gz", ".tgz"}:
            mode = "r:gz" if suffix in {".tar.gz", ".tgz"} else "r:"
            with tempfile.NamedTemporaryFile(suffix=suffix) as tmp:
                tmp.write(data)
                tmp.flush()
                with tarfile.open(tmp.name, mode) as archive:
                    if not archive.getmembers():
                        return "empty_tar_archive"
        elif suffix in {".xlsx", ".docx", ".pptx", ".odt", ".odp", ".ods"}:
            with tempfile.NamedTemporaryFile(suffix=suffix) as tmp:
                tmp.write(data)
                tmp.flush()
                path = Path(tmp.name)
                if suffix == ".xlsx":
                    load_workbook(path, read_only=True).close()
                else:
                    with zipfile.ZipFile(path) as archive:
                        names = set(archive.namelist())
                        if suffix == ".docx":
                            required = {"[Content_Types].xml", "_rels/.rels", "word/document.xml"}
                        elif suffix == ".pptx":
                            required = {"[Content_Types].xml", "_rels/.rels", "ppt/presentation.xml"}
                        else:
                            required = {"mimetype", "content.xml", "META-INF/manifest.xml"}
                        missing = sorted(required - names)
                        if missing:
                            return f"missing:{','.join(missing)}"
    except Exception as exc:
        return str(exc)
    return None


def _archive_member_format_errors(path: Path) -> list[tuple[str, str]]:
    if path.suffix.lower() != ".zip":
        return []
    checked_suffixes = {
        ".json",
        ".jpg",
        ".jpeg",
        ".png",
        ".mp3",
        ".wav",
        ".m4a",
        ".mp4",
        ".pdf",
        ".xlsx",
        ".docx",
        ".pptx",
        ".odt",
        ".odp",
        ".ods",
        ".tar",
        ".tar.gz",
        ".tgz",
    }
    errors: list[tuple[str, str]] = []
    try:
        with zipfile.ZipFile(path) as archive:
            for info in archive.infolist():
                if info.is_dir():
                    continue
                suffix = _member_extension(info.filename)
                if suffix not in checked_suffixes:
                    continue
                error = _document_bytes_format_error(info.filename, archive.read(info.filename))
                if error:
                    errors.append((info.filename, error))
    except Exception as exc:
        errors.append(("<archive>", str(exc)))
    return errors


def _upload_local_paths(task: dict[str, Any]) -> list[str]:
    paths: list[str] = []
    for section in ("setup", "cleanup"):
        for block in task.get(section, []) or []:
            for op in block.get("config", []) or []:
                if not isinstance(op, dict) or op.get("type") != "upload_file":
                    continue
                params = op.get("parameters", {}) or {}
                files = params.get("files") or []
                if isinstance(files, list):
                    for file_spec in files:
                        if isinstance(file_spec, dict) and file_spec.get("local_path"):
                            paths.append(str(file_spec["local_path"]))
                if params.get("local_path"):
                    paths.append(str(params["local_path"]))
    return paths


def _setup_uploads_by_runtime_path(task: dict[str, Any]) -> dict[tuple[str, str], str]:
    uploads: dict[tuple[str, str], str] = {}
    for block in task.get("setup", []) or []:
        if not isinstance(block, dict):
            continue
        device_id = str(block.get("device_id", ""))
        for op in block.get("config", []) or []:
            if not isinstance(op, dict) or op.get("type") != "upload_file":
                continue
            params = op.get("parameters", {}) or {}
            for file_spec in params.get("files") or []:
                if not isinstance(file_spec, dict):
                    continue
                local_path = file_spec.get("local_path")
                runtime_path = file_spec.get("path")
                if local_path and runtime_path:
                    uploads[(device_id, str(runtime_path))] = str(local_path)
            if params.get("local_path") and params.get("path"):
                uploads[(device_id, str(params["path"]))] = str(params["local_path"])
    return uploads


def _visible_text(task: dict[str, Any]) -> str:
    metadata = task.get("metadata", {}) or {}
    return _json_text(
        {
            "instruction": task.get("instruction"),
            "description": metadata.get("description"),
            "visible_sources": metadata.get("visible_sources"),
            "required_outputs": metadata.get("required_outputs"),
        }
    )


def _metadata_text(task: dict[str, Any]) -> str:
    text = _json_text(task.get("metadata", {}) or {})
    return re.sub(r'"local_path"\s*:\s*"[^"]+"', '"local_path":"<path>"', text)


def _ensure_app_categories(op: dict[str, Any]) -> set[str]:
    if op.get("type") != "ensure_app":
        return set()
    params = op.get("parameters", {}) or {}
    text = _json_text(params)
    categories: set[str] = set()
    for pattern, category in ANDROID_ENSURE_APP_CATEGORIES:
        if pattern.search(text):
            categories.add(category)
    return categories


def _android_helper_category(op_type: str) -> str | None:
    for pattern, category in ANDROID_HELPER_APP_CATEGORIES:
        if pattern.search(op_type):
            return category
    return None


def _iter_eval_items(task: dict[str, Any]) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for block in task.get("evaluation", []) or []:
        if not isinstance(block, dict):
            continue
        if "func" in block or "result" in block:
            items.append(block)
        for op in block.get("config", []) or []:
            if isinstance(op, dict):
                items.append({**op, "device_id": op.get("device_id", block.get("device_id"))})
    return items


def _expected_rule_value(evaluator: dict[str, Any]) -> str:
    expected = evaluator.get("expected", {}) or {}
    rules = expected.get("rules", {}) if isinstance(expected, dict) else {}
    if isinstance(rules, dict) and "expected" in rules:
        return str(rules.get("expected", ""))
    return ""


def _command_path_tokens(command: str) -> set[str]:
    return {
        token.strip("'\";")
        for token in re.findall(r"/(?:home/user|tmp|sdcard|storage/emulated/0)/[^\s;&|)]+", command)
        if token.strip("'\";")
    }


def _looks_like_file_value(value: str) -> bool:
    return bool(re.search(r"\.(?:csv|json|md|txt|m3u8?|mp3|m4a|wav|mp4|mov|png|jpe?g|pdf|odt|ods|odp|xlsx|zip)\b", value, re.I))


def _flag_shell_expected_transpositions(task: dict[str, Any], base: dict[str, Any], issues: dict[str, list[dict[str, Any]]]) -> None:
    eval_items = []
    for evaluator in _iter_eval_items(task):
        result = evaluator.get("result", {}) or {}
        if result.get("type") != "vm_command_line":
            continue
        command = result.get("command")
        if not isinstance(command, str):
            continue
        eval_items.append(
            {
                "device_id": evaluator.get("device_id"),
                "func": evaluator.get("func"),
                "command": command,
                "expected": _expected_rule_value(evaluator),
                "paths": _command_path_tokens(command),
            }
        )
    for first_line in eval_items:
        if "sed -n '1p'" not in first_line["command"] and 'sed -n "1p"' not in first_line["command"]:
            continue
        if not first_line["expected"].isdigit():
            continue
        for line_count in eval_items:
            if first_line is line_count or "wc -l" not in line_count["command"]:
                continue
            if not _looks_like_file_value(line_count["expected"]):
                continue
            if first_line["paths"] and line_count["paths"] and first_line["paths"].isdisjoint(line_count["paths"]):
                continue
            issues["shell_expected_transposed"].append(
                {
                    **base,
                    "device_id": first_line["device_id"],
                    "first_line_expected": first_line["expected"],
                    "line_count_expected": line_count["expected"],
                    "paths": sorted(first_line["paths"] | line_count["paths"]),
                }
            )


def check_task(
    *,
    repo_root: Path,
    round_name: str,
    sample: str,
    task_path: Path,
    task: dict[str, Any],
) -> dict[str, list[dict[str, Any]]]:
    issues: dict[str, list[dict[str, Any]]] = {key: [] for key in ISSUE_KEYS}
    task_id = str(task.get("id", ""))
    base = {"round": round_name, "sample": sample, "task_id": task_id, "task": str(task_path)}

    devices = {
        block.get("device_id")
        for section in ("setup", "evaluation", "cleanup")
        for block in (task.get(section, []) or [])
        if isinstance(block, dict) and block.get("device_id")
    }
    declared_devices = {
        device.get("id")
        for device in (task.get("devices", []) or [])
        if isinstance(device, dict) and device.get("id")
    }
    if len(devices) < 2:
        issues["single_topology"].append({**base, "devices": sorted(devices)})
    unused_declared = sorted(declared_devices - devices)
    if unused_declared:
        issues["declared_device_unused"].append({**base, "unused_devices": unused_declared})

    if _development_trace_pattern().search(_visible_text(task)):
        issues["visible_leak"].append(base)
    setup_trace = _setup_visible_trace(task)
    if setup_trace:
        issues["visible_leak"].append({**base, "trace": setup_trace})

    instruction = str(task.get("instruction", ""))
    if "/tmp/" in instruction or "/private/tmp" in instruction:
        issues["bad_linux_path"].append({**base, "instruction": instruction})

    metadata_surfaces = (task.get("metadata", {}) or {}).get("surfaces")
    if isinstance(metadata_surfaces, dict):
        issues["surface_schema"].append(base)
    elif isinstance(metadata_surfaces, list):
        allowed_surfaces = _allowed_surface_names(repo_root)
        invalid_surfaces = [
            surface
            for surface in metadata_surfaces
            if not isinstance(surface, str) or (allowed_surfaces and surface not in allowed_surfaces)
        ]
        if invalid_surfaces:
            issues["surface_schema"].append({**base, "invalid_surfaces": invalid_surfaces})

    if re.search(
        r"requires_stage|stage\d|pipeline|value_trace|engine_draft",
        _metadata_text(task),
        re.I,
    ):
        issues["metadata_trace"].append(base)

    full_text = _json_text(task)
    if re.search(r"Dummy MP3 content|fake audio|not a real audio|not real audio", full_text, re.I):
        issues["fake_audio_text"].append(base)

    if re.search(r"/home/oai/|file:///tmp/|/tmp/[^\"'\s]+\.html\b", full_text, re.I):
        issues["agent_inaccessible_path"].append(base)

    if "/data/media/0/Android/data/net.osmand" in full_text or "osmand/tracks" in full_text.lower():
        issues["old_osmand_path"].append(base)

    if re.search(r"/sdcard/AudioRecord(?:er|ings)(?:/|\\b)", full_text, re.I):
        issues["old_audio_recorder_path"].append(base)

    for block in task.get("setup", []) or []:
        seen_ensure = False
        ensured_categories: set[str] = set()
        for op in block.get("config", []) or []:
            if not isinstance(op, dict):
                continue
            op_type = op.get("type")
            if op_type == "ensure_app":
                seen_ensure = True
                ensured_categories.update(_ensure_app_categories(op))
            if op_type in {"androidworld_recipes_clear", "androidworld_recipe_add"} and not seen_ensure:
                issues["clear_before_ensure_setup"].append(
                    {**base, "device_id": block.get("device_id"), "op_type": op_type}
                )
            if isinstance(op_type, str):
                category = _android_helper_category(op_type)
                if category and category not in ensured_categories:
                    issues["android_helper_before_ensure"].append(
                        {
                            **base,
                            "device_id": block.get("device_id"),
                            "op_type": op_type,
                            "required_app_category": category,
                        }
                    )

    explicit_guard_pattern = re.compile(r"source sanity|guard only|not scored|non[- ]?scored", re.I)
    for evaluator in _iter_eval_items(task):
        if evaluator.get("enable_score_calc", True) is False:
            continue
        label_text = _json_text(
            {
                "name": evaluator.get("name"),
                "note": evaluator.get("note"),
                "notes": evaluator.get("notes"),
                "explanation": evaluator.get("explanation"),
            }
        )
        if explicit_guard_pattern.search(label_text):
            issues["explicit_guard_scored"].append(
                {**base, "device_id": evaluator.get("device_id"), "func": evaluator.get("func")}
            )

    supported_image_dimension_rules = {
        "count",
        "filenames",
        "exact_size",
        "all_square",
        "same_size",
        "min_width",
        "min_height",
        "max_width",
        "max_height",
    }
    for block in task.get("evaluation", []) or []:
        eval_items: list[dict[str, Any]] = []
        if isinstance(block, dict):
            eval_items.append(block)
            eval_items.extend(op for op in (block.get("config", []) or []) if isinstance(op, dict))
        for op in eval_items:
            if op.get("func") != "check_image_dimensions":
                continue
            rules = ((op.get("expected") or {}).get("rules") or {})
            unsupported = sorted(set(rules) - supported_image_dimension_rules)
            if unsupported:
                issues["unsupported_image_dimension_rule"].append(
                    {**base, "device_id": block.get("device_id"), "rules": rules, "unsupported": unsupported}
                )

    setup_uploads = _setup_uploads_by_runtime_path(task)
    for evaluator in _iter_eval_items(task):
        if evaluator.get("func") != "check_xlsx_cells":
            continue
        result = evaluator.get("result", {}) or {}
        runtime_path = result.get("path")
        device_id = str(evaluator.get("device_id", ""))
        rules = ((evaluator.get("expected") or {}).get("rules") or {})
        sheet = rules.get("sheet")
        if not runtime_path or not sheet:
            continue
        # New sheets are often valid task goals. As a hard gate, only block the
        # common generator fallback where the evaluator names "Sheet1" even
        # though the uploaded source workbook uses a natural sheet name.
        if sheet != "Sheet1":
            continue
        local_path = setup_uploads.get((device_id, str(runtime_path)))
        if not local_path:
            continue
        resolved = _resolve_repo_path(repo_root, local_path)
        if resolved.suffix.lower() != ".xlsx" or not resolved.exists():
            continue
        try:
            workbook = load_workbook(resolved, read_only=True, data_only=False)
            sheetnames = workbook.sheetnames
            workbook.close()
        except Exception as exc:
            issues["xlsx_sheet_missing"].append(
                {**base, "device_id": device_id, "path": str(runtime_path), "sheet": sheet, "error": str(exc)}
            )
            continue
        if sheet not in sheetnames:
            issues["xlsx_sheet_missing"].append(
                {
                    **base,
                    "device_id": device_id,
                    "path": str(runtime_path),
                    "sheet": sheet,
                    "available_sheets": sheetnames,
                }
            )

    _flag_shell_expected_transpositions(task, base, issues)

    for local_path in _upload_local_paths(task):
        resolved = _resolve_repo_path(repo_root, local_path)
        if not resolved.exists():
            issues["missing_upload"].append({**base, "local_path": local_path})
            continue
        if re.search(r"/(stage\d|generation_engine_probe|sample_\d{6})", local_path) and (
            local_path.endswith("/manifest.json") or local_path.endswith("/manifest.yaml")
        ):
            issues["development_manifest_upload"].append({**base, "local_path": local_path})
        suffix = resolved.suffix.lower()
        if suffix in {".mp3", ".wav", ".m4a", ".jpg", ".jpeg", ".png"}:
            description = _file_description(resolved)
            if not _media_format_ok(suffix, description):
                issues["fake_media"].append(
                    {**base, "local_path": local_path, "suffix": suffix, "file": description}
                )
        if suffix in {".xlsx", ".docx", ".odt", ".odp", ".ods", ".pdf"}:
            error = _document_format_error(resolved)
            if error:
                issues["bad_document_format"].append(
                    {**base, "local_path": local_path, "suffix": suffix, "error": error}
                )
        for member, error in _archive_member_format_errors(resolved):
            issues["bad_archive_member_format"].append(
                {**base, "local_path": local_path, "member": member, "error": error}
            )
        has_trace, trace = _asset_text_contains_trace(resolved)
        if has_trace:
            issues["visible_leak"].append({**base, "local_path": local_path, "trace": trace})

    return issues


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--manifest",
        default="runs/generation_engine_probe/stress_gpt5mini/R01_R35_selected_manifest.json",
    )
    parser.add_argument("--repo-root", default=".")
    parser.add_argument("--round", action="append", dest="rounds", help="Limit to one or more rounds, e.g. R01")
    parser.add_argument("--json", action="store_true", help="Print JSON details instead of a text summary")
    args = parser.parse_args()

    repo_root = Path(args.repo_root).resolve()
    manifest = Path(args.manifest)
    if not manifest.is_absolute():
        manifest = repo_root / manifest
    selected = _load_json(manifest)["selected"]
    rounds = set(args.rounds or [])
    if rounds:
        selected = [item for item in selected if item.get("round") in rounds]

    details: dict[str, list[dict[str, Any]]] = {key: [] for key in ISSUE_KEYS}
    for item in selected:
        task_path = Path(item["task"])
        task = _load_json(task_path)
        task_issues = check_task(
            repo_root=repo_root,
            round_name=str(item.get("round", "")),
            sample=str(item.get("sample", "")),
            task_path=task_path,
            task=task,
        )
        for key, rows in task_issues.items():
            details[key].extend(rows)

    if args.json:
        print(json.dumps({"selected": len(selected), "issues": details}, ensure_ascii=False, indent=2))
    else:
        print(f"selected {len(selected)}")
        for key in ISSUE_KEYS:
            print(f"{key} {len(details[key])}")
            for row in details[key][:5]:
                print(f"  {row}")

    return 1 if any(details.values()) else 0


if __name__ == "__main__":
    raise SystemExit(main())
