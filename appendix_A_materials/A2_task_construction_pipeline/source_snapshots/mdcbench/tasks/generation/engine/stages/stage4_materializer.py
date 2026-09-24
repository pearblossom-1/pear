from __future__ import annotations

import json
import math
import html
import re
from pathlib import Path, PurePosixPath
from io import BytesIO
import shutil
import sqlite3
import struct
import subprocess
import tempfile
from dataclasses import dataclass
from typing import Any, Callable
import wave
import zipfile
from xml.sax.saxutils import escape as xml_escape
import zlib


SUPPORTED_QUALITY_CHECKS = {
    "anchors_present",
    "audio_duration",
    "csv_header",
    "html_has_form",
    "html_bundle_files",
    "html_has_inputs",
    "image_dimensions",
    "json_parseable",
    "matches_visible_source",
    "media_duration",
    "minimum_rows",
    "no_deterministic_image_fallback",
    "nonempty_file",
    "nonempty_text",
    "office_text_anchors",
    "pdf_text_anchors",
    "submit_has_stable_hash",
    "zip_file_list",
}

FFMPEG_FIXTURE_TIMEOUT_S = 60


ImageGenerator = Callable[[dict[str, Any], Path], None]


@dataclass(frozen=True)
class MaterializationContext:
    image_generator: ImageGenerator | None = None
    deterministic_image_fallback: bool = False


class AssetMaterializer:
    materializer_name = ""
    asset_kinds: tuple[str, ...] = ()

    def write(
        self,
        target: Path,
        relative_path: str,
        content: str,
        asset: dict[str, Any],
        context: MaterializationContext,
    ) -> list[Path]:
        raise NotImplementedError


class PlainTextMaterializer(AssetMaterializer):
    materializer_name = "inline_text"
    asset_kinds = ("text_note", "text", "markdown", "receipt")

    def write(
        self,
        target: Path,
        relative_path: str,
        content: str,
        asset: dict[str, Any],
        context: MaterializationContext,
    ) -> list[Path]:
        target.write_text(content, encoding="utf-8")
        return [target]


class TabularCsvMaterializer(PlainTextMaterializer):
    materializer_name = "tabular_csv"
    asset_kinds = ("csv",)


class StructuredJsonMaterializer(PlainTextMaterializer):
    materializer_name = "structured_json"
    asset_kinds = ("json",)


class HtmlFormMaterializer(PlainTextMaterializer):
    materializer_name = "html_form"
    asset_kinds = ("html",)


class HtmlBundleMaterializer(AssetMaterializer):
    materializer_name = "html_bundle"
    asset_kinds = ("html_bundle",)

    def write(
        self,
        target: Path,
        relative_path: str,
        content: str,
        asset: dict[str, Any],
        context: MaterializationContext,
    ) -> list[Path]:
        spec = _json_spec(content)
        if spec:
            return _write_html_bundle(target, spec)
        target.write_text(content, encoding="utf-8")
        return [target]


class ZipArchiveMaterializer(AssetMaterializer):
    materializer_name = "zip_archive"
    asset_kinds = ("zip",)

    def write(
        self,
        target: Path,
        relative_path: str,
        content: str,
        asset: dict[str, Any],
        context: MaterializationContext,
    ) -> list[Path]:
        _write_zip(target, _json_spec(content))
        return [target]


class RasterImageMaterializer(AssetMaterializer):
    materializer_name = "raster_image"
    asset_kinds = ("image",)

    def write(
        self,
        target: Path,
        relative_path: str,
        content: str,
        asset: dict[str, Any],
        context: MaterializationContext,
    ) -> list[Path]:
        spec = _json_spec(content)
        width = int(spec.get("width", 640))
        height = int(spec.get("height", 360))
        label = _image_label_from_spec(spec)
        if spec.get("generation_backend") == "GPTImage2":
            if context.image_generator is not None:
                context.image_generator(spec, target)
                if not target.exists() or target.stat().st_size <= 0:
                    raise ValueError("GPTImage2 image generator did not write a non-empty file")
                return [target]
            if not context.deterministic_image_fallback:
                raise ValueError("GPTImage2 image generator is not configured")
        if target.suffix.lower() in {".jpg", ".jpeg"}:
            target.write_bytes(_jpeg_bytes(width, height, label))
            return [target]
        target.write_bytes(_png_bytes(width, height, label))
        return [target]


class AudioFixtureMaterializer(AssetMaterializer):
    materializer_name = "audio_fixture"
    asset_kinds = ("audio",)

    def write(
        self,
        target: Path,
        relative_path: str,
        content: str,
        asset: dict[str, Any],
        context: MaterializationContext,
    ) -> list[Path]:
        spec = _json_spec(content)
        duration_s = float(spec.get("duration_s", 1.0))
        sample_rate = int(spec.get("sample_rate", 16000))
        if target.suffix.lower() == ".m4a":
            _write_m4a(target, duration_s=duration_s, sample_rate=sample_rate)
            return [target]
        if target.suffix.lower() == ".mp3":
            _write_mp3(target, duration_s=duration_s, sample_rate=sample_rate)
            return [target]
        _write_wav(target, duration_s=duration_s, sample_rate=sample_rate)
        return [target]


class VideoFixtureMaterializer(AssetMaterializer):
    materializer_name = "video_fixture"
    asset_kinds = ("video",)

    def write(
        self,
        target: Path,
        relative_path: str,
        content: str,
        asset: dict[str, Any],
        context: MaterializationContext,
    ) -> list[Path]:
        _write_video(target, _json_spec(content))
        return [target]


def _image_label_from_spec(spec: dict[str, Any]) -> str:
    parts: list[str] = []
    for key in ("label", "image_prompt", "visual_requirements"):
        value = str(spec.get(key, "")).strip()
        if value:
            parts.append(value)
    elements = spec.get("elements")
    if isinstance(elements, list):
        for element in elements:
            if not isinstance(element, dict):
                continue
            for key in ("label", "text"):
                value = str(element.get(key, "")).strip()
                if value:
                    parts.append(value)
    anchors = spec.get("anchors")
    if isinstance(anchors, list):
        for anchor in anchors:
            value = str(anchor).strip()
            if value:
                parts.append(value)
    seen: set[str] = set()
    unique = []
    for part in parts:
        if part not in seen:
            seen.add(part)
            unique.append(part)
    return " | ".join(unique)[:240] or "visual fixture"


class DocumentTemplateMaterializer(AssetMaterializer):
    materializer_name = "document_template"
    asset_kinds = ("docx", "odp", "ods", "odt", "pdf", "xlsx")

    def write(
        self,
        target: Path,
        relative_path: str,
        content: str,
        asset: dict[str, Any],
        context: MaterializationContext,
    ) -> list[Path]:
        _write_document(target, _json_spec(content), str(asset.get("asset_kind", "")))
        return [target]


class CodeProjectMaterializer(AssetMaterializer):
    materializer_name = "code_project"
    asset_kinds = ("code_project",)

    def write(
        self,
        target: Path,
        relative_path: str,
        content: str,
        asset: dict[str, Any],
        context: MaterializationContext,
    ) -> list[Path]:
        spec = _json_spec(content)
        if target.suffix:
            raise ValueError("code_project relative_path must be a directory path, not a manifest/spec file")
        base_dir = target
        written: list[Path] = []
        for item in _list_value(spec.get("files")):
            if not isinstance(item, dict):
                continue
            relative = _safe_relative_path(str(item.get("path", "")))
            if relative is None:
                raise ValueError(f"unsafe code_project path {item.get('path', '')}")
            child = base_dir / relative.as_posix()
            child.parent.mkdir(parents=True, exist_ok=True)
            child.write_text(str(item.get("content", "")), encoding="utf-8")
            written.append(child)
        if not written:
            raise ValueError("code_project requires at least one file")
        return written


class GpxRouteMaterializer(AssetMaterializer):
    materializer_name = "gpx_route"
    asset_kinds = ("gpx",)

    def write(
        self,
        target: Path,
        relative_path: str,
        content: str,
        asset: dict[str, Any],
        context: MaterializationContext,
    ) -> list[Path]:
        _write_gpx(target, _json_spec(content))
        return [target]


class EmailMessageMaterializer(AssetMaterializer):
    materializer_name = "email_message"
    asset_kinds = ("email",)

    def write(
        self,
        target: Path,
        relative_path: str,
        content: str,
        asset: dict[str, Any],
        context: MaterializationContext,
    ) -> list[Path]:
        _write_email(target, _json_spec(content))
        return [target]


class PlaylistFileMaterializer(AssetMaterializer):
    materializer_name = "playlist_file"
    asset_kinds = ("playlist",)

    def write(
        self,
        target: Path,
        relative_path: str,
        content: str,
        asset: dict[str, Any],
        context: MaterializationContext,
    ) -> list[Path]:
        _write_playlist(target, _json_spec(content))
        return [target]


class SqliteDatabaseMaterializer(AssetMaterializer):
    materializer_name = "sqlite_database"
    asset_kinds = ("sqlite_db",)

    def write(
        self,
        target: Path,
        relative_path: str,
        content: str,
        asset: dict[str, Any],
        context: MaterializationContext,
    ) -> list[Path]:
        _write_sqlite(target, _json_spec(content))
        return [target]


MATERIALIZER_REGISTRY: dict[str, AssetMaterializer] = {
    materializer.materializer_name: materializer
    for materializer in (
        AudioFixtureMaterializer(),
        CodeProjectMaterializer(),
        DocumentTemplateMaterializer(),
        EmailMessageMaterializer(),
        GpxRouteMaterializer(),
        HtmlBundleMaterializer(),
        HtmlFormMaterializer(),
        PlainTextMaterializer(),
        PlaylistFileMaterializer(),
        RasterImageMaterializer(),
        SqliteDatabaseMaterializer(),
        StructuredJsonMaterializer(),
        TabularCsvMaterializer(),
        VideoFixtureMaterializer(),
        ZipArchiveMaterializer(),
    )
}

ASSET_KIND_MATERIALIZERS: dict[str, tuple[str, ...]] = {}
for materializer_name, materializer in MATERIALIZER_REGISTRY.items():
    for asset_kind in materializer.asset_kinds:
        ASSET_KIND_MATERIALIZERS.setdefault(asset_kind, tuple())
        ASSET_KIND_MATERIALIZERS[asset_kind] = ASSET_KIND_MATERIALIZERS[asset_kind] + (materializer_name,)


def materializer_for_asset(asset: dict[str, Any]) -> AssetMaterializer:
    materializer_name = str(asset.get("materializer", ""))
    asset_kind = str(asset.get("asset_kind", ""))
    materializer = MATERIALIZER_REGISTRY.get(materializer_name)
    if materializer is None:
        raise ValueError(f"unsupported materializer {materializer_name}")
    if asset_kind not in materializer.asset_kinds:
        raise ValueError(f"materializer {materializer_name} does not support asset_kind {asset_kind}")
    return materializer


def materialize_asset_plan(
    plan: dict[str, Any],
    *,
    output_dir: str | Path,
    image_generator: ImageGenerator | None = None,
    deterministic_image_fallback: bool = False,
) -> dict[str, Any]:
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    for child in output_path.iterdir():
        if child.is_dir():
            shutil.rmtree(child)
        else:
            child.unlink()
    written_files: list[dict[str, Any]] = []
    errors: list[str] = []
    context = MaterializationContext(
        image_generator=image_generator,
        deterministic_image_fallback=deterministic_image_fallback,
    )
    assets = {
        str(asset.get("asset_id")): asset
        for asset in plan.get("asset_manifest", [])
        if isinstance(asset, dict) and asset.get("asset_id")
    }

    for file_item in plan.get("external_files", []):
        if not isinstance(file_item, dict):
            errors.append("external_files item must be an object")
            continue
        relative_path = str(file_item.get("relative_path", ""))
        content = _content_to_text(file_item.get("content"))
        if not relative_path:
            errors.append("external_files item missing relative_path")
            continue
        if not content.strip():
            errors.append(f"{relative_path} content must be non-empty for materialization")
            continue
        target = output_path / relative_path
        try:
            target.relative_to(output_path)
        except ValueError:
            errors.append(f"{relative_path} escapes output_dir")
            continue
        target.parent.mkdir(parents=True, exist_ok=True)
        asset = assets.get(str(file_item.get("asset_id", "")), {})
        try:
            for written in _write_materialized_file(target, relative_path, content, asset, context):
                written_files.append(
                    {
                        "asset_id": file_item.get("asset_id"),
                        "relative_path": written.relative_to(output_path).as_posix(),
                        "bytes": written.stat().st_size,
                    }
                )
        except (ValueError, OSError, subprocess.SubprocessError) as exc:
            errors.append(f"{relative_path} materialization failed: {exc}")

    quality_report = run_asset_quality_checks(plan, materialized_dir=output_path)
    errors.extend(quality_report["errors"])
    return {
        "passed": not errors,
        "written_files": written_files,
        "checks": quality_report["checks"],
        "errors": errors,
    }


def run_asset_quality_checks(plan: dict[str, Any], *, materialized_dir: str | Path | None = None) -> dict[str, Any]:
    files_by_asset = {
        str(file_item.get("asset_id")): file_item
        for file_item in plan.get("external_files", [])
        if isinstance(file_item, dict)
    }
    checks: dict[str, dict[str, bool]] = {}
    errors: list[str] = []

    for asset in plan.get("asset_manifest", []):
        if not isinstance(asset, dict):
            continue
        asset_id = str(asset.get("asset_id", ""))
        if not asset_id:
            continue
        file_item = files_by_asset.get(asset_id)
        content = ""
        if isinstance(file_item, dict):
            content = _content_to_text(file_item.get("content", ""))
        materialized_path = None
        if materialized_dir is not None and isinstance(file_item, dict):
            materialized_path = Path(materialized_dir) / str(file_item.get("relative_path", ""))
        asset_checks: dict[str, bool] = {}
        for check in asset.get("quality_checks", []):
            result = _run_single_check(str(check), asset, content, materialized_path)
            asset_checks[str(check)] = result
            if not result:
                errors.append(_quality_error(str(check), asset, content))
        checks[asset_id] = asset_checks

    return {"passed": not errors, "checks": checks, "errors": errors}


def _write_materialized_file(
    target: Path,
    relative_path: str,
    content: str,
    asset: dict[str, Any],
    context: MaterializationContext,
) -> list[Path]:
    return materializer_for_asset(asset).write(target, relative_path, content, asset, context)


def _content_to_text(content: Any) -> str:
    if isinstance(content, str):
        return content
    if isinstance(content, (dict, list)):
        return json.dumps(content, ensure_ascii=False)
    if content is None:
        return ""
    return str(content)


def _run_single_check(check: str, asset: dict[str, Any], content: str, materialized_path: Path | None = None) -> bool:
    if check in {"nonempty_text", "nonempty_file"}:
        if materialized_path is not None and materialized_path.exists():
            return materialized_path.stat().st_size > 0
        return bool(content.strip() or str(asset.get("content", "")).strip())
    if check == "matches_visible_source":
        return bool(str(asset.get("content", "")).strip())
    if check == "anchors_present":
        haystack = f"{content}\n{asset.get('content', '')}\n{_materialized_text(materialized_path)}"
        return all(_anchor_present(str(anchor), haystack) for anchor in _asset_anchor_values(asset))
    if check == "csv_header":
        first_line = next((line for line in content.splitlines() if line.strip()), "")
        return "," in first_line and len([cell for cell in first_line.split(",") if cell.strip()]) >= 2
    if check == "minimum_rows":
        return _minimum_rows_ok(content)
    if check == "html_has_form":
        lower = content.lower()
        return "<form" in lower and "</form>" in lower
    if check == "submit_has_stable_hash":
        lower = content.lower()
        return "submit" in lower and "#" in content and "location" in lower
    if check == "html_has_inputs":
        lower = content.lower()
        return "<input" in lower or "<textarea" in lower or "<select" in lower
    if check == "html_bundle_files":
        return _html_bundle_files_ok(content, materialized_path)
    if check == "json_parseable":
        try:
            json.loads(content)
        except json.JSONDecodeError:
            return False
        return True
    if check == "zip_file_list":
        return _zip_file_list_ok(content, asset, materialized_path)
    if check == "image_dimensions":
        return _image_dimensions_ok(content, materialized_path)
    if check == "no_deterministic_image_fallback":
        return not _looks_like_deterministic_image_fallback(materialized_path)
    if check == "audio_duration":
        return _audio_duration_ok(content, materialized_path)
    if check == "media_duration":
        return _media_duration_ok(content, materialized_path)
    if check == "office_text_anchors":
        return _office_text_anchors_ok(asset, content, materialized_path)
    if check == "pdf_text_anchors":
        return _pdf_text_anchors_ok(asset, content, materialized_path)
    return False


def _quality_error(check: str, asset: dict[str, Any], content: str) -> str:
    asset_id = str(asset.get("asset_id", "<unknown>"))
    if check == "anchors_present":
        haystack = f"{content}\n{asset.get('content', '')}"
        missing = [str(anchor) for anchor in _asset_anchor_values(asset) if not _anchor_present(str(anchor), haystack)]
        return f"{asset_id} missing anchor values: {', '.join(missing)}"
    if check == "submit_has_stable_hash":
        anchors = ", ".join(str(anchor) for anchor in _asset_anchor_values(asset))
        return f"{asset_id} missing stable submit hash for anchors: {anchors}"
    if check not in SUPPORTED_QUALITY_CHECKS:
        return f"{asset_id} uses unsupported quality check {check}"
    return f"{asset_id} failed quality check {check}"


def _materialized_text(materialized_path: Path | None) -> str:
    if materialized_path is None or not materialized_path.exists():
        return ""
    if materialized_path.is_file():
        return materialized_path.read_text(encoding="utf-8", errors="ignore")
    chunks: list[str] = []
    for child in sorted(materialized_path.rglob("*")):
        if child.is_file():
            chunks.append(child.read_text(encoding="utf-8", errors="ignore"))
    return "\n".join(chunks)


def _json_spec(content: str) -> dict[str, Any]:
    try:
        value = json.loads(content)
    except json.JSONDecodeError:
        return {}
    return value if isinstance(value, dict) else {}


def _list_value(value: Any) -> list[Any]:
    return value if isinstance(value, list) else []


def _asset_anchor_values(asset: dict[str, Any]) -> list[Any]:
    return _list_value(asset.get("expected_anchor_values"))


def _anchor_present(anchor: str, haystack: str) -> bool:
    if anchor in haystack:
        return True
    grep_match = re.search(r"""grep\s+-q\s+(['"])(.+?)\1""", anchor)
    return bool(grep_match and grep_match.group(2) in haystack)


def _minimum_rows_ok(content: str) -> bool:
    spec = _json_spec(content)
    sheets = spec.get("sheets") if isinstance(spec.get("sheets"), list) else []
    for sheet in sheets:
        if not isinstance(sheet, dict):
            continue
        rows = sheet.get("rows")
        if isinstance(rows, list) and max(0, len([row for row in rows if isinstance(row, list)]) - 1) >= 8:
            return True
    tables = spec.get("tables") if isinstance(spec.get("tables"), list) else []
    for table in tables:
        if not isinstance(table, dict):
            continue
        rows = table.get("rows")
        if isinstance(rows, list) and len([row for row in rows if isinstance(row, (list, dict))]) >= 8:
            return True
    files = spec.get("files") if isinstance(spec.get("files"), list) else []
    for item in files:
        if not isinstance(item, dict):
            continue
        file_content = str(item.get("content", ""))
        if _minimum_rows_ok(file_content):
            return True
    nonempty_lines = [line for line in content.splitlines() if line.strip()]
    return max(0, len(nonempty_lines) - 1) >= 8


def _write_zip(target: Path, spec: dict[str, Any]) -> None:
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for item in _list_value(spec.get("files")):
            if not isinstance(item, dict):
                continue
            name = str(item.get("path", "")).strip("/")
            if not name or ".." in PurePosixPath(name).parts:
                continue
            archive.writestr(name, _zip_member_bytes(name, item.get("content", "")))


def _zip_member_bytes(name: str, content: Any) -> bytes:
    suffix = PurePosixPath(name).suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".tsv", ".html", ".htm", ".xml", ".log", ".ini"}:
        return _content_to_text(content).encode("utf-8")
    if suffix == ".json":
        if isinstance(content, (dict, list)):
            return json.dumps(content, ensure_ascii=False, indent=2).encode("utf-8")
        text = _content_to_text(content)
        json.loads(text)
        return text.encode("utf-8")
    if suffix in {".png", ".jpg", ".jpeg"}:
        spec = _zip_document_spec(name, content)
        width = int(spec.get("width", 640))
        height = int(spec.get("height", 360))
        label = _image_label_from_spec(spec)
        if suffix in {".jpg", ".jpeg"}:
            return _jpeg_bytes(width, height, label)
        return _png_bytes(width, height, label)
    if suffix in {".wav", ".m4a", ".mp3", ".mp4", ".pdf", ".xlsx", ".docx", ".pptx", ".odt", ".ods", ".odp"}:
        return _materialized_member_bytes(name, suffix, content)
    return _content_to_text(content).encode("utf-8")


def _zip_document_spec(name: str, content: Any) -> dict[str, Any]:
    if isinstance(content, dict):
        return content
    text = _content_to_text(content).strip()
    spec = _json_spec(text)
    if spec:
        return spec
    stem = PurePosixPath(name).stem.replace("_", " ").replace("-", " ").strip() or "Document"
    if not text:
        text = stem
    return {"title": stem.title(), "paragraphs": [text], "anchors": [text]}


def _materialized_member_bytes(name: str, suffix: str, content: Any) -> bytes:
    spec = _zip_document_spec(name, content)
    with tempfile.TemporaryDirectory(prefix="mdcbench_zip_member_") as temp_dir:
        target = Path(temp_dir) / PurePosixPath(name).name
        if suffix == ".wav":
            _write_wav(
                target,
                duration_s=float(spec.get("duration_s", 1.0)),
                sample_rate=int(spec.get("sample_rate", 16000)),
            )
        elif suffix == ".m4a":
            _write_m4a(
                target,
                duration_s=float(spec.get("duration_s", 1.0)),
                sample_rate=int(spec.get("sample_rate", 16000)),
            )
        elif suffix == ".mp3":
            _write_mp3(
                target,
                duration_s=float(spec.get("duration_s", 1.0)),
                sample_rate=int(spec.get("sample_rate", 16000)),
            )
        elif suffix == ".mp4":
            _write_video(target, spec)
        elif suffix == ".pptx":
            _write_pptx(target, _document_lines(spec))
        elif suffix in {".pdf", ".xlsx", ".docx", ".odt", ".ods", ".odp"}:
            _write_document(target, spec, suffix.lstrip("."))
        else:
            raise ValueError(f"unsupported zip member format {suffix}")
        return target.read_bytes()


def _write_html_bundle(target: Path, spec: dict[str, Any]) -> list[Path]:
    files = spec.get("files", [])
    if not isinstance(files, list) or not files:
        raise ValueError("html_bundle requires non-empty files")
    base_dir = target.parent
    written: list[Path] = []
    for item in files:
        if not isinstance(item, dict):
            continue
        relative = _safe_relative_path(str(item.get("path", "")))
        if relative is None:
            raise ValueError(f"unsafe html_bundle path {item.get('path', '')}")
        child = base_dir / relative.as_posix()
        child.parent.mkdir(parents=True, exist_ok=True)
        child.write_text(str(item.get("content", "")), encoding="utf-8")
        written.append(child)
    if not written:
        raise ValueError("html_bundle did not write any files")
    return written


def _write_wav(target: Path, *, duration_s: float, sample_rate: int) -> None:
    frame_count = max(1, int(duration_s * sample_rate))
    amplitude = 8000
    frequency = 440
    with wave.open(str(target), "wb") as wav:
        wav.setnchannels(1)
        wav.setsampwidth(2)
        wav.setframerate(sample_rate)
        frames = bytearray()
        for index in range(frame_count):
            sample = int(amplitude * math.sin(2 * math.pi * frequency * index / sample_rate))
            frames.extend(struct.pack("<h", sample))
        wav.writeframes(bytes(frames))


def _write_m4a(target: Path, *, duration_s: float, sample_rate: int) -> None:
    ffmpeg = shutil.which("ffmpeg")
    if not ffmpeg:
        raise ValueError("ffmpeg is required to generate m4a audio fixtures")
    command = [
        ffmpeg,
        "-nostdin",
        "-y",
        "-hide_banner",
        "-loglevel",
        "error",
        "-f",
        "lavfi",
        "-i",
        f"sine=frequency=440:sample_rate={sample_rate}:duration={duration_s}",
        "-c:a",
        "aac",
        "-b:a",
        "96k",
        str(target),
    ]
    subprocess.run(command, check=True, stdin=subprocess.DEVNULL, timeout=FFMPEG_FIXTURE_TIMEOUT_S)


def _write_mp3(target: Path, *, duration_s: float, sample_rate: int) -> None:
    ffmpeg = shutil.which("ffmpeg")
    if not ffmpeg:
        raise ValueError("ffmpeg is required to generate mp3 audio fixtures")
    command = [
        ffmpeg,
        "-nostdin",
        "-y",
        "-hide_banner",
        "-loglevel",
        "error",
        "-f",
        "lavfi",
        "-i",
        f"sine=frequency=440:sample_rate={sample_rate}:duration={duration_s}",
        "-q:a",
        "6",
        str(target),
    ]
    subprocess.run(command, check=True, stdin=subprocess.DEVNULL, timeout=FFMPEG_FIXTURE_TIMEOUT_S)


def _write_video(target: Path, spec: dict[str, Any]) -> None:
    ffmpeg = shutil.which("ffmpeg")
    if not ffmpeg:
        raise ValueError("ffmpeg is required for video_fixture materialization")
    duration_s = max(0.5, float(spec.get("duration_s", 2.0)))
    width = max(16, int(spec.get("width", 320)))
    height = max(16, int(spec.get("height", 180)))
    if width % 2:
        width += 1
    if height % 2:
        height += 1
    command = [
        ffmpeg,
        "-nostdin",
        "-y",
        "-hide_banner",
        "-loglevel",
        "error",
        "-f",
        "lavfi",
        "-i",
        f"color=c=0x336699:s={width}x{height}:d={duration_s}",
        "-pix_fmt",
        "yuv420p",
        str(target),
    ]
    subprocess.run(command, check=True, stdin=subprocess.DEVNULL, timeout=FFMPEG_FIXTURE_TIMEOUT_S)


def _write_document(target: Path, spec: dict[str, Any], asset_kind: str) -> None:
    if asset_kind == "pdf":
        _write_pdf(target, _document_lines(spec))
        return
    if asset_kind == "xlsx":
        _write_xlsx(target, spec)
        return
    if asset_kind == "docx":
        _write_docx(target, _document_lines(spec))
        return
    if asset_kind in {"odt", "ods", "odp"}:
        _write_odf(target, spec, asset_kind)
        return
    target.write_text("\n".join(_document_lines(spec)) + "\n", encoding="utf-8")


def _document_lines(spec: dict[str, Any]) -> list[str]:
    lines: list[str] = []
    title = spec.get("title")
    if title:
        lines.append(str(title))
    for paragraph in _list_value(spec.get("paragraphs")):
        lines.append(str(paragraph))
    for slide in _list_value(spec.get("slides")):
        if not isinstance(slide, dict):
            continue
        if slide.get("title"):
            lines.append(str(slide.get("title")))
        for paragraph in _list_value(slide.get("paragraphs")):
            lines.append(str(paragraph))
    for sheet in _list_value(spec.get("sheets")):
        if not isinstance(sheet, dict):
            continue
        if sheet.get("name"):
            lines.append(str(sheet.get("name")))
        for row in _list_value(sheet.get("rows")):
            if isinstance(row, list):
                lines.append(" | ".join(str(cell) for cell in row))
    for table in _list_value(spec.get("tables")):
        if not isinstance(table, dict):
            continue
        table_name = table.get("table_name", table.get("name"))
        if table_name:
            lines.append(str(table_name))
        columns = [str(column) for column in _list_value(table.get("columns")) if column is not None]
        if columns:
            lines.append(" | ".join(columns))
        for row in _list_value(table.get("rows")):
            if isinstance(row, list):
                lines.append(" | ".join(str(cell) for cell in row))
            elif isinstance(row, dict):
                if columns:
                    lines.append(" | ".join(str(row.get(column, "")) for column in columns))
                else:
                    lines.append(" | ".join(f"{key}: {value}" for key, value in row.items()))
    for anchor in _list_value(spec.get("anchors")):
        if str(anchor) not in "\n".join(lines):
            lines.append(str(anchor))
    for anchor in _list_value(spec.get("text_anchors")):
        if str(anchor) not in "\n".join(lines):
            lines.append(str(anchor))
    return lines or ["Document fixture"]


def _write_pdf(target: Path, lines: list[str]) -> None:
    def pdf_text(value: str) -> str:
        return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    stream_lines = ["BT", "/F1 12 Tf", "72 740 Td", "16 TL"]
    for line in lines:
        stream_lines.append(f"({pdf_text(line)}) Tj")
        stream_lines.append("T*")
    stream_lines.append("ET")
    stream = "\n".join(stream_lines).encode("utf-8")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("ascii"))
        output.extend(obj)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii"))
    target.write_bytes(bytes(output))


def _write_odf(target: Path, spec: dict[str, Any], asset_kind: str) -> None:
    mimetype = {
        "odt": "application/vnd.oasis.opendocument.text",
        "ods": "application/vnd.oasis.opendocument.spreadsheet",
        "odp": "application/vnd.oasis.opendocument.presentation",
    }[asset_kind]
    body_tag = {
        "odt": "office:text",
        "ods": "office:spreadsheet",
        "odp": "office:presentation",
    }[asset_kind]
    body_content = _odf_body_content(spec, asset_kind)
    content_xml = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
        "<office:document-content "
        "xmlns:office=\"urn:oasis:names:tc:opendocument:xmlns:office:1.0\" "
        "xmlns:text=\"urn:oasis:names:tc:opendocument:xmlns:text:1.0\" "
        "xmlns:table=\"urn:oasis:names:tc:opendocument:xmlns:table:1.0\" "
        "xmlns:draw=\"urn:oasis:names:tc:opendocument:xmlns:drawing:1.0\" "
        "xmlns:presentation=\"urn:oasis:names:tc:opendocument:xmlns:presentation:1.0\" "
        "xmlns:style=\"urn:oasis:names:tc:opendocument:xmlns:style:1.0\">"
        "<office:scripts/>"
        "<office:automatic-styles/>"
        f"<office:body><{body_tag}>{body_content}</{body_tag}></office:body>"
        "</office:document-content>"
    )
    styles_xml = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
        "<office:document-styles "
        "xmlns:office=\"urn:oasis:names:tc:opendocument:xmlns:office:1.0\" "
        "xmlns:style=\"urn:oasis:names:tc:opendocument:xmlns:style:1.0\" "
        "xmlns:text=\"urn:oasis:names:tc:opendocument:xmlns:text:1.0\">"
        "<office:styles><style:default-style style:family=\"paragraph\"/></office:styles>"
        "</office:document-styles>"
    )
    meta_xml = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
        "<office:document-meta "
        "xmlns:office=\"urn:oasis:names:tc:opendocument:xmlns:office:1.0\" "
        "xmlns:meta=\"urn:oasis:names:tc:opendocument:xmlns:meta:1.0\">"
        "<office:meta><meta:generator>OpenDocument Generator</meta:generator></office:meta>"
        "</office:document-meta>"
    )
    manifest_xml = _odf_manifest_xml(mimetype)
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", mimetype, compress_type=zipfile.ZIP_STORED)
        archive.writestr("content.xml", content_xml)
        archive.writestr("styles.xml", styles_xml)
        archive.writestr("meta.xml", meta_xml)
        archive.writestr("META-INF/manifest.xml", manifest_xml)


def _odf_body_content(spec: dict[str, Any], asset_kind: str) -> str:
    if asset_kind == "ods":
        sheets = [sheet for sheet in _list_value(spec.get("sheets")) if isinstance(sheet, dict)]
        if not sheets:
            rows = [[line] for line in _document_lines(spec)]
            sheets = [{"name": "Sheet1", "rows": rows}]
        return "".join(_odf_table_xml(sheet, index) for index, sheet in enumerate(sheets, start=1))
    if asset_kind == "odp":
        slides = [slide for slide in _list_value(spec.get("slides")) if isinstance(slide, dict)]
        if not slides:
            slides = [{"title": line, "paragraphs": []} for line in _document_lines(spec)]
        return "".join(_odf_slide_xml(slide, index) for index, slide in enumerate(slides, start=1))
    return "".join(f"<text:p>{xml_escape(line)}</text:p>" for line in _document_lines(spec))


def _odf_table_xml(sheet: dict[str, Any], index: int) -> str:
    name = _xml_attr_escape(str(sheet.get("name") or f"Sheet{index}"))
    rows = sheet.get("rows")
    if not isinstance(rows, list) or not rows:
        rows = [[line] for line in _document_lines({"paragraphs": [name]})]
    row_xml = []
    for row in rows:
        if not isinstance(row, list):
            continue
        cells = "".join(
            "<table:table-cell office:value-type=\"string\">"
            f"<text:p>{xml_escape(str(cell))}</text:p>"
            "</table:table-cell>"
            for cell in row
        )
        row_xml.append(f"<table:table-row>{cells}</table:table-row>")
    return f"<table:table table:name=\"{name}\">{''.join(row_xml)}</table:table>"


def _odf_slide_xml(slide: dict[str, Any], index: int) -> str:
    name = _xml_attr_escape(str(slide.get("name") or slide.get("title") or f"Slide {index}"))
    lines: list[str] = []
    if slide.get("title"):
        lines.append(str(slide.get("title")))
    lines.extend(str(paragraph) for paragraph in _list_value(slide.get("paragraphs")))
    if not lines:
        lines = [name]
    paragraphs = "".join(f"<text:p>{xml_escape(line)}</text:p>" for line in lines)
    return (
        f"<draw:page draw:name=\"{name}\" presentation:presentation-page-layout-name=\"AL1T0\">"
        f"{paragraphs}"
        "</draw:page>"
    )


def _odf_manifest_xml(mimetype: str) -> str:
    return (
        "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
        "<manifest:manifest "
        "xmlns:manifest=\"urn:oasis:names:tc:opendocument:xmlns:manifest:1.0\" "
        "manifest:version=\"1.2\">"
        f"<manifest:file-entry manifest:full-path=\"/\" manifest:media-type=\"{xml_escape(mimetype)}\"/>"
        "<manifest:file-entry manifest:full-path=\"content.xml\" manifest:media-type=\"text/xml\"/>"
        "<manifest:file-entry manifest:full-path=\"styles.xml\" manifest:media-type=\"text/xml\"/>"
        "<manifest:file-entry manifest:full-path=\"meta.xml\" manifest:media-type=\"text/xml\"/>"
        "</manifest:manifest>"
    )


def _write_docx(target: Path, lines: list[str]) -> None:
    body = "".join(f"<w:p><w:r><w:t>{xml_escape(line)}</w:t></w:r></w:p>" for line in lines)
    document = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>"
        "<w:document xmlns:w=\"http://schemas.openxmlformats.org/wordprocessingml/2006/main\">"
        f"<w:body>{body}</w:body></w:document>"
    )
    root_rels = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
        "<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\">"
        "<Relationship Id=\"rId1\" Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument\" Target=\"word/document.xml\"/>"
        "</Relationships>"
    )
    content_types = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
        "<Types xmlns=\"http://schemas.openxmlformats.org/package/2006/content-types\">"
        "<Default Extension=\"rels\" ContentType=\"application/vnd.openxmlformats-package.relationships+xml\"/>"
        "<Default Extension=\"xml\" ContentType=\"application/xml\"/>"
        "<Override PartName=\"/word/document.xml\" ContentType=\"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml\"/>"
        "</Types>"
    )
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", root_rels)
        archive.writestr("word/document.xml", document)


def _write_pptx(target: Path, lines: list[str]) -> None:
    title = lines[0] if lines else "Slide"
    body = "\n".join(lines[1:] or lines or ["Slide content"])
    presentation = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>"
        "<p:presentation xmlns:p=\"http://schemas.openxmlformats.org/presentationml/2006/main\" "
        "xmlns:r=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships\">"
        "<p:sldIdLst><p:sldId id=\"256\" r:id=\"rId1\"/></p:sldIdLst>"
        "</p:presentation>"
    )
    slide = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>"
        "<p:sld xmlns:p=\"http://schemas.openxmlformats.org/presentationml/2006/main\" "
        "xmlns:a=\"http://schemas.openxmlformats.org/drawingml/2006/main\">"
        "<p:cSld><p:spTree><p:nvGrpSpPr/><p:grpSpPr/>"
        "<p:sp><p:nvSpPr/><p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/>"
        f"<a:p><a:r><a:t>{xml_escape(title)}</a:t></a:r></a:p>"
        f"<a:p><a:r><a:t>{xml_escape(body)}</a:t></a:r></a:p>"
        "</p:txBody></p:sp></p:spTree></p:cSld></p:sld>"
    )
    root_rels = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
        "<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\">"
        "<Relationship Id=\"rId1\" Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument\" Target=\"ppt/presentation.xml\"/>"
        "</Relationships>"
    )
    presentation_rels = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
        "<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\">"
        "<Relationship Id=\"rId1\" Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide\" Target=\"slides/slide1.xml\"/>"
        "</Relationships>"
    )
    content_types = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
        "<Types xmlns=\"http://schemas.openxmlformats.org/package/2006/content-types\">"
        "<Default Extension=\"rels\" ContentType=\"application/vnd.openxmlformats-package.relationships+xml\"/>"
        "<Default Extension=\"xml\" ContentType=\"application/xml\"/>"
        "<Override PartName=\"/ppt/presentation.xml\" ContentType=\"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml\"/>"
        "<Override PartName=\"/ppt/slides/slide1.xml\" ContentType=\"application/vnd.openxmlformats-officedocument.presentationml.slide+xml\"/>"
        "</Types>"
    )
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", root_rels)
        archive.writestr("ppt/presentation.xml", presentation)
        archive.writestr("ppt/_rels/presentation.xml.rels", presentation_rels)
        archive.writestr("ppt/slides/slide1.xml", slide)


def _write_xlsx(target: Path, spec: dict[str, Any]) -> None:
    sheets = spec.get("sheets") if isinstance(spec.get("sheets"), list) else []
    strings: list[str] = []
    string_index: dict[str, int] = {}
    normalized_sheets: list[tuple[str, list[Any]]] = []
    for index, sheet in enumerate(sheets, start=1):
        if not isinstance(sheet, dict):
            continue
        rows = sheet.get("rows", [])
        if not isinstance(rows, list) or not rows:
            continue
        name = str(sheet.get("name") or f"Sheet{index}")
        normalized_sheets.append((name, rows))
    if not normalized_sheets:
        normalized_sheets = [("Sheet1", [[line] for line in _document_lines(spec)])]

    worksheets: list[str] = []
    for _sheet_name, rows in normalized_sheets:
        cells_xml: list[str] = []
        for row_index, row in enumerate(rows, start=1):
            if not isinstance(row, list):
                continue
            cell_parts = []
            for col_index, value in enumerate(row, start=1):
                text = str(value)
                if text not in string_index:
                    string_index[text] = len(strings)
                    strings.append(text)
                cell_ref = f"{_excel_col(col_index)}{row_index}"
                cell_parts.append(f"<c r=\"{cell_ref}\" t=\"s\"><v>{string_index[text]}</v></c>")
            cells_xml.append(f"<row r=\"{row_index}\">{''.join(cell_parts)}</row>")
        worksheets.append(
            "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
            "<worksheet xmlns=\"http://schemas.openxmlformats.org/spreadsheetml/2006/main\">"
            f"<sheetData>{''.join(cells_xml)}</sheetData></worksheet>"
        )
    shared_strings = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
        f"<sst xmlns=\"http://schemas.openxmlformats.org/spreadsheetml/2006/main\" count=\"{len(strings)}\" uniqueCount=\"{len(strings)}\">"
        + "".join(f"<si><t>{xml_escape(value)}</t></si>" for value in strings)
        + "</sst>"
    )
    sheet_entries = "".join(
        f"<sheet name=\"{_xml_attr_escape(name)}\" sheetId=\"{index}\" r:id=\"rId{index}\"/>"
        for index, (name, _rows) in enumerate(normalized_sheets, start=1)
    )
    workbook = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
        "<workbook xmlns=\"http://schemas.openxmlformats.org/spreadsheetml/2006/main\" "
        "xmlns:r=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships\">"
        f"<sheets>{sheet_entries}</sheets></workbook>"
    )
    worksheet_relationships = "".join(
        f"<Relationship Id=\"rId{index}\" Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet\" Target=\"worksheets/sheet{index}.xml\"/>"
        for index in range(1, len(normalized_sheets) + 1)
    )
    shared_string_rid = len(normalized_sheets) + 1
    rels = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
        "<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\">"
        f"{worksheet_relationships}"
        f"<Relationship Id=\"rId{shared_string_rid}\" Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings\" Target=\"sharedStrings.xml\"/>"
        "</Relationships>"
    )
    root_rels = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
        "<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\">"
        "<Relationship Id=\"rId1\" Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument\" Target=\"xl/workbook.xml\"/>"
        "</Relationships>"
    )
    worksheet_content_types = "".join(
        f"<Override PartName=\"/xl/worksheets/sheet{index}.xml\" ContentType=\"application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml\"/>"
        for index in range(1, len(normalized_sheets) + 1)
    )
    content_types = "".join(
        [
            "<?xml version=\"1.0\" encoding=\"UTF-8\"?>",
            "<Types xmlns=\"http://schemas.openxmlformats.org/package/2006/content-types\">",
            "<Default Extension=\"rels\" ContentType=\"application/vnd.openxmlformats-package.relationships+xml\"/>",
            "<Default Extension=\"xml\" ContentType=\"application/xml\"/>",
            "<Override PartName=\"/xl/workbook.xml\" ContentType=\"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml\"/>",
            "<Override PartName=\"/xl/sharedStrings.xml\" ContentType=\"application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml\"/>",
            worksheet_content_types,
            "</Types>",
        ]
    )
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", root_rels)
        archive.writestr("xl/workbook.xml", workbook)
        archive.writestr("xl/_rels/workbook.xml.rels", rels)
        for index, worksheet in enumerate(worksheets, start=1):
            archive.writestr(f"xl/worksheets/sheet{index}.xml", worksheet)
        archive.writestr("xl/sharedStrings.xml", shared_strings)


def _excel_col(index: int) -> str:
    name = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        name = chr(65 + remainder) + name
    return name or "A"


def _xml_attr_escape(value: str) -> str:
    return xml_escape(value, {'"': "&quot;"})


def _write_gpx(target: Path, spec: dict[str, Any]) -> None:
    name = str(spec.get("name", "Waypoint"))
    lat = str(spec.get("lat", spec.get("latitude", "0")))
    lon = str(spec.get("lon", spec.get("longitude", "0")))
    desc = str(spec.get("desc", spec.get("description", "")))
    gpx = (
        "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
        "<gpx version=\"1.1\" creator=\"OsmAnd\" xmlns=\"http://www.topografix.com/GPX/1/1\">"
        f"<wpt lat=\"{xml_escape(lat)}\" lon=\"{xml_escape(lon)}\">"
        f"<name>{xml_escape(name)}</name><desc>{xml_escape(desc)}</desc>"
        "</wpt></gpx>"
    )
    target.write_text(gpx, encoding="utf-8")


def _write_email(target: Path, spec: dict[str, Any]) -> None:
    if target.suffix.lower() == ".mbox":
        messages = spec.get("messages")
        if not isinstance(messages, list) or not messages:
            messages = [spec]
        chunks: list[str] = []
        for index, message in enumerate(messages):
            if not isinstance(message, dict):
                continue
            sender = str(message.get("from", "sender@example.test"))
            recipient = str(message.get("to", spec.get("to", "recipient@example.test")))
            subject = str(message.get("subject", "Request"))
            body = str(message.get("body", ""))
            date = str(message.get("date", f"Sat Jun 07 08:{index:02d}:00 2026"))
            chunks.append(
                f"From {sender} {date}\n"
                f"From: {sender}\n"
                f"To: {recipient}\n"
                f"Subject: {subject}\n\n"
                f"{body}\n"
            )
        target.write_text("\n".join(chunks), encoding="utf-8")
        return
    sender = str(spec.get("from", "sender@example.test"))
    recipient = str(spec.get("to", "recipient@example.test"))
    subject = str(spec.get("subject", "Request"))
    body = str(spec.get("body", ""))
    target.write_text(
        f"From: {sender}\nTo: {recipient}\nSubject: {subject}\nMIME-Version: 1.0\nContent-Type: text/plain; charset=utf-8\n\n{body}\n",
        encoding="utf-8",
    )


def _write_playlist(target: Path, spec: dict[str, Any]) -> None:
    tracks = [str(track) for track in _list_value(spec.get("tracks"))]
    target.write_text("#EXTM3U\n" + "\n".join(tracks) + "\n", encoding="utf-8")


def _write_sqlite(target: Path, spec: dict[str, Any]) -> None:
    if target.exists():
        target.unlink()
    with sqlite3.connect(target) as conn:
        for table in _list_value(spec.get("tables")):
            if not isinstance(table, dict):
                continue
            name = _sql_identifier(str(table.get("name", "")))
            columns = table.get("columns", [])
            if not name or not isinstance(columns, list) or not columns:
                continue
            column_defs = []
            column_names = []
            for column in columns:
                if not isinstance(column, list) or len(column) < 2:
                    continue
                col_name = _sql_identifier(str(column[0]))
                col_type = str(column[1]).upper()
                if not col_name:
                    continue
                if col_type not in {"TEXT", "INTEGER", "REAL", "BLOB"}:
                    col_type = "TEXT"
                column_defs.append(f'"{col_name}" {col_type}')
                column_names.append(col_name)
            if not column_defs:
                continue
            conn.execute(f'CREATE TABLE "{name}" ({", ".join(column_defs)})')
            placeholders = ", ".join("?" for _ in column_names)
            quoted = ", ".join(f'"{column}"' for column in column_names)
            for row in table.get("rows", []):
                if isinstance(row, list):
                    conn.execute(f'INSERT INTO "{name}" ({quoted}) VALUES ({placeholders})', row[: len(column_names)])
        conn.commit()


def _sql_identifier(value: str) -> str:
    return "".join(char for char in value if char.isalnum() or char == "_")


def _safe_relative_path(value: str) -> PurePosixPath | None:
    path = PurePosixPath(value.strip())
    if not path.as_posix() or path.is_absolute() or ".." in path.parts:
        return None
    return path


def _png_bytes(width: int, height: int, label: str) -> bytes:
    width = max(1, width)
    height = max(1, height)
    rows = bytearray()
    label_seed = sum(label.encode("utf-8")) % 255
    for y in range(height):
        rows.append(0)
        for x in range(width):
            r = (x * 3 + label_seed) % 256
            g = (y * 5 + 80) % 256
            b = ((x + y) * 2 + 120) % 256
            rows.extend((r, g, b))
    chunks = [
        _png_chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)),
        _png_chunk(b"tEXt", f"Description\x00{label}".encode("utf-8")),
        _png_chunk(b"IDAT", zlib.compress(bytes(rows), level=6)),
        _png_chunk(b"IEND", b""),
    ]
    return b"\x89PNG\r\n\x1a\n" + b"".join(chunks)


def _png_chunk(kind: bytes, data: bytes) -> bytes:
    return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)


def _jpeg_bytes(width: int, height: int, label: str) -> bytes:
    from PIL import Image, ImageDraw

    width = max(1, width)
    height = max(1, height)
    label_seed = sum(label.encode("utf-8")) % 255
    image = Image.new("RGB", (width, height), (230, 235, 242))
    pixels = image.load()
    for y in range(height):
        for x in range(width):
            pixels[x, y] = (
                (x * 3 + label_seed) % 256,
                (y * 5 + 80) % 256,
                ((x + y) * 2 + 120) % 256,
            )
    draw = ImageDraw.Draw(image)
    draw.rectangle((8, 8, max(20, width - 8), min(height - 8, 44)), fill=(255, 255, 255), outline=(40, 60, 90))
    draw.text((14, 18), label[:80], fill=(20, 30, 45))
    output = BytesIO()
    image.save(output, format="JPEG", quality=90)
    return output.getvalue()


def _zip_file_list_ok(content: str, asset: dict[str, Any], materialized_path: Path | None) -> bool:
    spec = _json_spec(content)
    expected = {
        str(item.get("path", "")).strip("/")
        for item in _list_value(spec.get("files"))
        if isinstance(item, dict) and item.get("path")
    }
    if materialized_path is not None and materialized_path.exists():
        try:
            with zipfile.ZipFile(materialized_path) as archive:
                names = set(archive.namelist())
        except zipfile.BadZipFile:
            return False
        return bool(expected) and expected.issubset(names)
    haystack = f"{content}\n{asset.get('content', '')}"
    return all(name in haystack for name in expected)


def _html_bundle_files_ok(content: str, materialized_path: Path | None) -> bool:
    spec = _json_spec(content)
    if not spec:
        return bool(materialized_path is not None and materialized_path.exists())
    expected = [
        _safe_relative_path(str(item.get("path", "")))
        for item in _list_value(spec.get("files"))
        if isinstance(item, dict) and item.get("path")
    ]
    if not expected or any(path is None for path in expected):
        return False
    if materialized_path is None:
        return True
    base_dir = materialized_path.parent
    return all((base_dir / path.as_posix()).exists() for path in expected if path is not None)


def _image_dimensions_ok(content: str, materialized_path: Path | None) -> bool:
    spec = _json_spec(content)
    expected_width = int(spec.get("width", 0))
    expected_height = int(spec.get("height", 0))
    if materialized_path is None or not materialized_path.exists() or expected_width <= 0 or expected_height <= 0:
        return False
    data = materialized_path.read_bytes()
    if data.startswith(b"\x89PNG\r\n\x1a\n") and len(data) >= 24:
        return int.from_bytes(data[16:20], "big") == expected_width and int.from_bytes(data[20:24], "big") == expected_height
    try:
        from PIL import Image

        with Image.open(materialized_path) as image:
            return image.size == (expected_width, expected_height)
    except Exception:
        return False


def _looks_like_deterministic_image_fallback(materialized_path: Path | None) -> bool:
    if materialized_path is None or not materialized_path.exists():
        return False
    try:
        from PIL import Image

        with Image.open(materialized_path) as image:
            rgb = image.convert("RGB")
    except Exception:
        return False

    width, height = rgb.size
    if width <= 0 or height <= 0:
        return False
    pixels = rgb.load()
    xs = [int((index + 0.5) * width / 18) for index in range(18)]
    ys = [int((index + 0.5) * height / 18) for index in range(18)]
    # The JPEG fallback draws a small label strip near the top; skip it so
    # the signature check compares the deterministic gradient body.
    ys = [y for y in ys if y >= min(height - 1, max(45, height // 4))] or ys
    best_mean_error = math.inf
    for label_seed in range(255):
        total_error = 0
        channel_count = 0
        for y in ys:
            if y >= height:
                continue
            for x in xs:
                if x >= width:
                    continue
                red, green, blue = pixels[x, y]
                total_error += abs(red - ((x * 3 + label_seed) % 256))
                total_error += abs(green - ((y * 5 + 80) % 256))
                total_error += abs(blue - (((x + y) * 2 + 120) % 256))
                channel_count += 3
        if channel_count:
            best_mean_error = min(best_mean_error, total_error / channel_count)
    return best_mean_error < 15


def _audio_duration_ok(content: str, materialized_path: Path | None) -> bool:
    spec = _json_spec(content)
    expected = float(spec.get("duration_s", 0))
    if materialized_path is None or not materialized_path.exists() or expected <= 0:
        return False
    if materialized_path.suffix.lower() == ".m4a":
        actual = _ffprobe_duration(materialized_path)
        return actual is not None and actual >= expected * 0.9
    try:
        with wave.open(str(materialized_path), "rb") as wav:
            actual = wav.getnframes() / float(wav.getframerate())
    except (wave.Error, EOFError):
        return False
    return actual >= expected * 0.9


def _media_duration_ok(content: str, materialized_path: Path | None) -> bool:
    spec = _json_spec(content)
    expected = float(spec.get("duration_s", 0))
    if materialized_path is None or not materialized_path.exists() or expected <= 0:
        return False
    if materialized_path.suffix.lower() == ".wav":
        return _audio_duration_ok(content, materialized_path)
    actual = _ffprobe_duration(materialized_path)
    return actual is not None and actual >= expected * 0.9


def _ffprobe_duration(materialized_path: Path) -> float | None:
    ffprobe = shutil.which("ffprobe")
    if not ffprobe:
        return None
    try:
        output = subprocess.check_output(
            [
                ffprobe,
                "-v",
                "error",
                "-show_entries",
                "format=duration",
                "-of",
                "default=noprint_wrappers=1:nokey=1",
                str(materialized_path),
            ],
            text=True,
        ).strip()
        return float(output)
    except (subprocess.SubprocessError, ValueError):
        return None


def _office_text_anchors_ok(asset: dict[str, Any], content: str, materialized_path: Path | None) -> bool:
    anchors = _text_content_anchor_values(asset)
    if not anchors:
        return False
    haystack = content
    if materialized_path is not None and materialized_path.exists():
        if materialized_path.suffix.lower() == ".xlsx" and not _xlsx_openable(materialized_path):
            return False
        try:
            with zipfile.ZipFile(materialized_path) as archive:
                xml_parts = []
                for name in archive.namelist():
                    if name.endswith(".xml"):
                        xml_parts.append(archive.read(name).decode("utf-8", errors="ignore"))
                haystack = html.unescape("\n".join(xml_parts))
        except zipfile.BadZipFile:
            haystack = materialized_path.read_text(encoding="utf-8", errors="ignore")
    return all(anchor in haystack for anchor in anchors)


def _xlsx_openable(path: Path) -> bool:
    try:
        from openpyxl import load_workbook

        load_workbook(path, read_only=True, data_only=True).close()
        return True
    except Exception:
        return False


def _pdf_text_anchors_ok(asset: dict[str, Any], content: str, materialized_path: Path | None) -> bool:
    anchors = _text_content_anchor_values(asset)
    if not anchors:
        return False
    haystack = content
    if materialized_path is not None and materialized_path.exists():
        data = materialized_path.read_bytes()
        haystack = "\n".join(
            (
                data.decode("utf-8", errors="ignore"),
                data.decode("latin-1", errors="ignore"),
            )
        )
        haystack = _pdf_literal_unescape(haystack)
    return all(anchor in haystack for anchor in anchors)


def _pdf_literal_unescape(text: str) -> str:
    return (
        text.replace("\\\\", "\\")
        .replace("\\(", "(")
        .replace("\\)", ")")
    )


def _text_content_anchor_values(asset: dict[str, Any]) -> list[str]:
    return [
        str(anchor)
        for anchor in _asset_anchor_values(asset)
        if not str(anchor).strip().startswith("/")
    ]
